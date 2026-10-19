"""Read user-provided datasets with bounded shape checks."""

import csv
import re
from decimal import Decimal
from pathlib import Path

import pandas as pd
from config import Config
from utils.validators import DatasetLimitExceeded, dataset_limit_error_from_shape

_INT_RE = re.compile(r"[+-]?\d+\Z")


def list_excel_sheets(filepath: str | Path) -> list[str]:
    """Return worksheet names of an Excel file (empty on any failure)."""
    try:
        with pd.ExcelFile(filepath) as book:
            return [str(name) for name in book.sheet_names]
    except Exception:
        return []


def describe_parse_error(error: Exception) -> str:
    """Turn a parser exception into one actionable sentence for the user."""
    name = type(error).__name__
    text = str(error).strip().splitlines()[0] if str(error).strip() else ""
    if isinstance(error, DatasetLimitExceeded):
        return str(error)
    if isinstance(error, UnicodeDecodeError):
        return (
            "This file is not valid UTF-8 text. Choose the Latin-1 (Western) "
            "encoding and upload it again."
        )
    if isinstance(error, FileNotFoundError) or "No such file" in text:
        return "The uploaded file is no longer on disk. Please upload it again."
    if "Worksheet" in text or "sheet" in text.lower():
        return f"The worksheet could not be read: {text}"
    if "BadZipFile" in name or "zip" in text.lower() or "excel file format" in text.lower():
        return "This .xlsx file looks damaged. Re-save it as .xlsx or export it as CSV."
    if "ParserError" in name or "delimiter" in text.lower() or "tokenizing" in text.lower():
        return (
            f"The rows do not line up as CSV: {text}. Check the delimiter "
            "(comma, semicolon, tab) and the encoding, then try again."
        )
    if not text:
        return f"We could not read this file ({name})."
    return f"We could not read this file ({name}: {text})."


def load_dataset_with_error(
    filepath: str | Path,
    encoding: str | None = None,
    delimiter: str | None = None,
    sheet: str | None = None,
) -> tuple[pd.DataFrame | None, str]:
    """Return a parsed dataset or an actionable reason for failure."""
    try:
        return _read_dataset(filepath, encoding=encoding, delimiter=delimiter, sheet=sheet), ""
    except Exception as error:
        return None, describe_parse_error(error)


def load_dataset(
    filepath: str | Path,
    encoding: str | None = None,
    delimiter: str | None = None,
    sheet: str | None = None,
) -> pd.DataFrame | None:
    """Read a CSV, Excel or internal pickle dataset, returning None on failure."""
    try:
        return _read_dataset(filepath, encoding=encoding, delimiter=delimiter, sheet=sheet)
    except Exception:
        return None


def _read_dataset(
    filepath: str | Path,
    encoding: str | None = None,
    delimiter: str | None = None,
    sheet: str | None = None,
) -> pd.DataFrame:
    """Read a dataset, raising on failure so the caller can explain why."""
    filepath = Path(filepath)
    if not filepath.exists() or not filepath.is_file():
        raise FileNotFoundError(str(filepath))

    suffix = filepath.suffix.lower()
    if suffix == ".pkl":
        return pd.read_pickle(filepath)
    if suffix == ".csv":
        kwargs: dict = {}
        if encoding:
            kwargs["encoding"] = encoding
        if delimiter and delimiter != "auto":
            kwargs["delimiter"] = delimiter
        else:
            kwargs["sep"] = None
            kwargs["engine"] = "python"
        effective_encoding = _preflight_csv_limits(
            filepath, encoding=encoding, delimiter=delimiter
        )
        if effective_encoding and not encoding:
            kwargs["encoding"] = effective_encoding
        try:
            df = pd.read_csv(filepath, **kwargs)
        except UnicodeDecodeError:
            if encoding:
                raise
            fallback = {key: value for key, value in kwargs.items() if key != "encoding"}
            df = pd.read_csv(filepath, encoding="latin1", **fallback)
            kwargs["encoding"] = "latin1"
        repair_columns = [
            column for column in df.columns
            if pd.api.types.is_numeric_dtype(df[column])
        ]
        if not repair_columns:
            return df
        raw_kwargs = dict(kwargs)
        raw_kwargs.update(
            dtype=str,
            keep_default_na=False,
            usecols=repair_columns,
        )
        df_raw = pd.read_csv(filepath, **raw_kwargs)
        _repair_typed_columns(df, df_raw)
        return df
    if suffix in (".xls", ".xlsx"):
        _preflight_excel_limits(filepath, sheet=sheet)
        if sheet:
            return pd.read_excel(filepath, sheet_name=sheet)
        return pd.read_excel(filepath)
    raise ValueError(f"Unsupported file type '{suffix}'.")


def _preflight_csv_limits(
    filepath: Path,
    *,
    encoding: str | None,
    delimiter: str | None,
) -> str | None:
    """Count CSV records without materializing them into a DataFrame."""
    encodings = [encoding] if encoding else ["utf-8", "latin1"]
    for candidate_encoding in encodings:
        try:
            with filepath.open("r", encoding=candidate_encoding, newline="") as handle:
                sample = handle.read(8192)
                handle.seek(0)
                csv_delimiter = delimiter if delimiter and delimiter != "auto" else None
                if not csv_delimiter:
                    try:
                        csv_delimiter = csv.Sniffer().sniff(
                            sample, delimiters=",;\t"
                        ).delimiter
                    except csv.Error:
                        csv_delimiter = ","
                reader = csv.reader(handle, delimiter=csv_delimiter)
                header = next(reader, None)
                if header is None:
                    return candidate_encoding
                columns = len(header)
                _raise_if_over_limit(0, columns)
                rows = 0
                for row in reader:
                    if not row or all(not value.strip() for value in row):
                        continue
                    rows += 1
                    _raise_if_over_limit(rows, columns)
                return candidate_encoding
        except UnicodeDecodeError:
            if encoding:
                raise
    return encoding


def _preflight_excel_limits(filepath: Path, *, sheet: str | None) -> None:
    """Inspect workbook dimensions before pandas materializes the dataset."""
    if filepath.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(filepath, read_only=True, data_only=True)
            try:
                if sheet:
                    if sheet not in workbook.sheetnames:
                        raise ValueError(f"Worksheet '{sheet}' not found.")
                    worksheet = workbook[sheet]
                else:
                    worksheet = workbook.active
                columns = int(worksheet.max_column or 0)
                rows = max(int(worksheet.max_row or 0) - 1, 0)
                _raise_if_over_limit(rows, columns)
            finally:
                workbook.close()
            return
        except ImportError:
            pass

    probe = pd.read_excel(filepath, sheet_name=sheet, nrows=1)
    columns = int(probe.shape[1])
    _raise_if_over_limit(0, columns)
    if columns <= 0:
        return
    sample_rows = min(
        Config.MAX_ROWS + 1,
        (Config.MAX_CELLS // columns) + 1,
    )
    sample = pd.read_excel(filepath, sheet_name=sheet, nrows=sample_rows)
    _raise_if_over_limit(len(sample), columns)


def _raise_if_over_limit(rows: int, columns: int) -> None:
    """Raise as soon as a bounded read proves a configured limit is exceeded."""
    message = dataset_limit_error_from_shape(rows, columns)
    if message:
        raise DatasetLimitExceeded(message)


def _repair_typed_columns(df: pd.DataFrame, df_raw: pd.DataFrame) -> None:
    """Preserve leading-zero identifiers and integers wider than float precision."""
    for column in df.columns:
        if column not in df_raw.columns:
            continue
        raw = df_raw[column].astype("string").str.strip()
        raw = raw.mask((raw == "") | df[column].isna())
        values = raw.dropna()
        if values.empty or not bool(values.str.fullmatch(_INT_RE).all()):
            continue
        if bool(values.str.fullmatch(r"[+-]?0\d+").any()):
            df[column] = raw
            continue

        numbers = [Decimal(value) for value in values]
        if max(abs(number) for number in numbers) <= 2**53:
            continue
        if max(abs(number) for number in numbers) <= 2**63 - 1:
            from features.cleaning.operations import _parse_integer

            df[column] = pd.array(
                [_parse_integer(value) if pd.notna(value) else pd.NA for value in raw],
                dtype="Int64",
            )
        else:
            df[column] = raw
