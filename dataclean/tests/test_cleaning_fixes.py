"""Comprehensive unit and integration tests verifying all functional inconsistency fixes."""

import io
import json
import os
import sys
import time
from pathlib import Path
from uuid import uuid4

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import create_app
from config import Config
from features.cleaning import operations as cleaning_service
from features.dataset import analysis as analysis_service
from features.upload import dataset_importer
from workspace import dataset_exporter, file_store
from utils.validators import dataset_limit_error

app = create_app()
app.config["TESTING"] = True
messy_csv = Path(__file__).parent / "sample_messy_data.csv"

failures = []


def check(name, condition, extra=""):
    status = "PASS" if condition else "FAIL"
    if not condition:
        failures.append(name)
    print(f"[{status}] {name}" + (f" -> {extra}" if extra and not condition else ""))


def current_csrf(client) -> str:
    """Read the token a freshly rendered form would submit."""
    with client.session_transaction() as active_session:
        return active_session.get("csrf_token", "test-token")


# --------------------------------------------------------------------------
# 1. Test Category Standardization NaN bug
# --------------------------------------------------------------------------
df_cat = pd.DataFrame({"gender": ["male", np.nan, "female", np.nan, "male"]})
cleaned_cat, count = cleaning_service.standardize_categories(df_cat, "gender", {"male": "Male"})
check(
    "standardize_categories does not count NaNs as changed",
    count == 2,
    f"expected 2 changes, got {count}",
)
check(
    "standardize_categories correctly replaces values",
    cleaned_cat["gender"].iloc[0] == "Male" and pd.isna(cleaned_cat["gender"].iloc[1]),
)

# --------------------------------------------------------------------------
# 1b. Safe integer conversion: no silent rounding, reject keeps the column
#     with examples reported, coerce moves invalids to missing, missing is
#     preserved and excluded from counts.
# --------------------------------------------------------------------------
df_int = pd.DataFrame({"score": ["10", "12.5", "bad", "", None, "3"]})
rejected_df, reject_report = cleaning_service.convert_column_type(df_int, "score", "integer", "reject")
check(
    "integer conversion rejected when fractional/invalid present",
    not reject_report.applied and list(rejected_df["score"].dropna().astype(str)) == list(df_int["score"].dropna().astype(str)),
)
check(
    "integer conversion reports invalid examples",
    reject_report.invalid_count == 3 and "12.5" in reject_report.invalid_examples,
    f"invalid={reject_report.invalid_count}, examples={reject_report.invalid_examples}",
)
coerced_df, coerce_report = cleaning_service.convert_column_type(df_int, "score", "integer", "coerce_to_missing")
check(
    "coerced integer conversion keeps whole numbers exactly",
    coerced_df["score"].tolist() == [10, pd.NA, pd.NA, pd.NA, pd.NA, 3],
    f"got {coerced_df['score'].tolist()}",
)
check(
    "coerced integer conversion counts changes and missing preserved",
    # Empty string is a non-integer value, so it is coerced to missing too.
    coerce_report.changed_values == 5 and coerce_report.coerced_count == 3,
    f"changed={coerce_report.changed_values}, coerced={coerce_report.coerced_count}",
)

df_round = pd.DataFrame({"n": ["12.5"]})
cleaned_round, round_report = cleaning_service.convert_column_type(df_round, "n", "integer", "reject")
check(
    "integer conversion does not silently round decimals",
    cleaned_round["n"].iloc[0] == "12.5" and not round_report.applied,
)

# --------------------------------------------------------------------------
# 2. Test Text Case Standardization does NOT strip whitespace
# --------------------------------------------------------------------------
df_case = pd.DataFrame({"dept": [" cs ", "cs", "CS"]})
cleaned_case, case_count = cleaning_service.standardize_text_case(df_case, "dept", "title")
check(
    "standardize_text_case preserves whitespace without stripping",
    cleaned_case["dept"].iloc[0] == " Cs ",
    f"got {cleaned_case['dept'].iloc[0]!r}",
)
check(
    "standardize_text_case accurate change count",
    case_count == 3,  # ' cs ' -> ' Cs ', 'cs' -> 'Cs', 'CS' -> 'Cs'
    f"got {case_count}",
)


# --------------------------------------------------------------------------
# 3. Test Numeric Custom Fill Preserves Numeric Type
# --------------------------------------------------------------------------
df_num = pd.DataFrame({"age": [20.0, np.nan, 30.0]})
cleaned_num, num_count = cleaning_service.handle_missing_values(df_num, "age", "custom", "25")
check("custom missing fill on numeric column count", num_count == 1)
check(
    "custom missing fill retains numeric dtype",
    pd.api.types.is_numeric_dtype(cleaned_num["age"]),
    f"dtype is {cleaned_num['age'].dtype}",
)
check("custom missing filled with numeric 25", cleaned_num["age"].iloc[1] == 25.0)

df_bool = pd.DataFrame({"active": ["yes", "No", "1", "0", "unknown"]})
cleaned_bool, bool_report = cleaning_service.convert_column_type(
    df_bool, "active", "boolean", "coerce_to_missing"
)
check("boolean conversion counts changed values", bool_report.changed_values == 5)
check("boolean conversion reports invalid values", bool_report.coerced_count == 1)
check(
    "boolean conversion marks invalid values missing",
    pd.isna(cleaned_bool["active"].iloc[4]),
)

# --------------------------------------------------------------------------
# 4. Test Invalid Date Cleaning
# --------------------------------------------------------------------------
df_date = pd.DataFrame({"date_col": ["2023-01-01", "not-a-date", "2023-05-10"]})
cleaned_date_coerce, n_coerce = cleaning_service.clean_invalid_dates(df_date, "date_col", "coerce")
check("clean_invalid_dates coerce count", n_coerce == 1)
check("clean_invalid_dates coerce sets NA", pd.isna(cleaned_date_coerce["date_col"].iloc[1]))

cleaned_date_remove, n_remove = cleaning_service.clean_invalid_dates(df_date, "date_col", "remove")
check("clean_invalid_dates remove count", n_remove == 1)
check("clean_invalid_dates removes row", len(cleaned_date_remove) == 2)

# --------------------------------------------------------------------------
# 5. Test Date Hints False Positives Prevention
# --------------------------------------------------------------------------
df_false_dates = pd.DataFrame({
    "candidate_name": ["Alice", "Bob", "Charlie", "David"],
    "centimeter_measurement": [12.5, 14.2, 10.1, 15.0],
})
basic_issues = analysis_service.detect_basic_issues(df_false_dates)
date_issues = [i for i in basic_issues if i["issue"] == "invalid_date"]
check(
    "date hint does not flag candidate or centimeter as date",
    len(date_issues) == 0,
    f"detected unexpected date issues: {date_issues}",
)

# --------------------------------------------------------------------------
# 6. Test Specific Column Drop
# --------------------------------------------------------------------------
df_drop = pd.DataFrame({"a": [1, 2], "b": [3, 4], "c": [5, 6]})
cleaned_drop, n_dropped = cleaning_service.drop_columns(df_drop, ["b", "c"])
check("drop_columns removes specified columns", n_dropped == 2 and list(cleaned_drop.columns) == ["a"])

# --------------------------------------------------------------------------
# 6b. Duplicate handling: honored mode, custom key columns, missing keys.
# --------------------------------------------------------------------------
df_dup = pd.DataFrame({"id": [1, 1, 2, 2, 3], "val": [10, 11, 20, 21, 30]})
first_df, n_first = cleaning_service.remove_duplicate_rows(df_dup, "first", ["id"])
check(
    "dedup keep first by key",
    n_first == 2 and first_df["val"].tolist() == [10, 20, 30],
    f"n={n_first}, kept={first_df['val'].tolist()}",
)
last_df, n_last = cleaning_service.remove_duplicate_rows(df_dup, "last", ["id"])
check(
    "dedup keep last by key",
    n_last == 2 and last_df["val"].tolist() == [11, 21, 30],
    f"n={n_last}, kept={last_df['val'].tolist()}",
)
all_df, n_all = cleaning_service.remove_duplicate_rows(df_dup, "all", ["id"])
check(
    "dedup remove all occurrences by key",
    n_all == 4 and all_df["val"].tolist() == [30],
    f"n={n_all}, kept={all_df['val'].tolist()}",
)
full_df, n_full = cleaning_service.remove_duplicate_rows(df_dup)
check(
    "dedup defaults to the entire row (all distinct, 0 removed)",
    n_full == 0 and len(full_df) == 5,
)

df_miss = pd.DataFrame({"id": [1, None, 1, None], "val": ["a", "b", "c", "d"]})
miss_df, n_miss = cleaning_service.remove_duplicate_rows(df_miss, "first", ["id"])
check(
    "dedup keys with missing values keep first occurrence",
    n_miss == 2 and miss_df["val"].tolist() == ["a", "b"],
    f"n={n_miss}, kept={miss_df['val'].tolist()}",
)
all_miss_df, n_all_miss = cleaning_service.remove_duplicate_rows(df_miss, "all", ["id"])
check(
    "dedup keys with missing values remove all occurrences",
    n_all_miss == 4 and all_miss_df.empty,
    f"n={n_all_miss}, kept={len(all_miss_df)}",
)

# --------------------------------------------------------------------------
# 6c. Missing-value handling: whitespace-only recognized, tokens opt-in,
#     custom value type validation, detection/cleaning agreement, tie note.
# --------------------------------------------------------------------------
df_miss_tok = pd.DataFrame({"score": ["20", None, "N/A", "  ", "30", "n/A"]})
missing_counts = analysis_service.get_missing_values(df_miss_tok)
check(
    "analysis counts whitespace-only and tokens as missing",
    missing_counts.get("score") == 4,
    f"got {missing_counts}",
)
missing_tokens = analysis_service.get_missing_token_counts(df_miss_tok)
check(
    "analysis reports token breakdown per column",
    missing_tokens.get("score") == 2,
    f"got {missing_tokens}",
)

kept_df, n_kept = cleaning_service.handle_missing_values(df_miss_tok, "score", "remove")
check(
    "tokens left alone without explicit selection",
    n_kept == 2 and "N/A" in kept_df["score"].tolist() and "  " not in kept_df["score"].tolist(),
    f"n={n_kept}, kept={kept_df['score'].tolist()}",
)
filled_df, n_filled = cleaning_service.handle_missing_values(
    df_miss_tok, "score", "remove", None, set(Config.MISSING_TOKENS)
)
check(
    "token opt-in cleans exactly what detection flagged",
    n_filled == 4 and filled_df["score"].tolist() == ["20", "30"],
    f"n={n_filled}, kept={filled_df['score'].tolist()}",
)

df_custom = pd.DataFrame({"age": [30, None, 40]})
try:
    cleaning_service.handle_missing_values(df_custom, "age", "custom", "abc")
    check("invalid custom value rejected", False, "no error raised")
except ValueError as exc:
    check("invalid custom value raises clear error", "not a number" in str(exc), str(exc))

df_intfill = pd.DataFrame({"count": pd.array([1, None, 3], dtype="Int64")})
try:
    cleaning_service.handle_missing_values(df_intfill, "count", "custom", 2.5)
    check("fractional custom fill on integer column rejected", False, "no error raised")
except ValueError as exc:
    check(
        "fractional custom fill rejected with clear error",
        "whole number" in str(exc),
        str(exc),
    )
filled_custom, n_custom = cleaning_service.handle_missing_values(df_intfill, "count", "custom", "5")
check(
    "valid custom fill applies and keeps integer dtype",
    n_custom == 1 and filled_custom["count"].tolist() == [1, 5, 3],
    f"n={n_custom}, got={filled_custom['count'].tolist()}",
)

df_blank = pd.DataFrame({"blank": ["", "   ", ""], "ok": [1, 2, 3]})
empty_issues = [i for i in analysis_service.detect_basic_issues(df_blank) if i["issue"] == "empty"]
check(
    "whitespace-only column detected as empty",
    len(empty_issues) == 1 and empty_issues[0]["column"] == "blank",
    f"got {empty_issues}",
)
dropped_blank, n_dropped_blank = cleaning_service.remove_empty_columns(df_blank)
check(
    "whitespace-only column removed by cleaning",
    n_dropped_blank == 1 and list(dropped_blank.columns) == ["ok"],
)

df_tie = pd.DataFrame({"team": ["A", "B", None]})
cleaned_tie, hist_tie = cleaning_service.apply_cleaning_operations(
    df_tie, [{"type": "missing", "column": "team", "method": "mode"}]
)
check(
    "mode tie reported explicitly in history",
    any("tied" in line for line in hist_tie),
    str(hist_tie),
)

# --------------------------------------------------------------------------
# 6d. Import/export fidelity: leading-zero IDs, oversized integers,
#     type-preserving round trips, CSV/Excel controls.
#     (Uses the Flask app only for the config context dataset exporters expect.)
# --------------------------------------------------------------------------
app4 = create_app()
app4.config["TESTING"] = True

csv_bytes = (
    b"ID,Amount,Label\n"
    b"00123,9007199254740993,alpha\n"
    b"00124,42,beta\n"
    b"00125,,gamma\n"
)
semi_bytes = "Vornamen;Sta_dtte\nJ\u00fcrgen;K\u00f6ln\n".encode("latin1")

with app4.app_context():
    from features.upload import dataset_importer as importer4
    from workspace import dataset_exporter as exporter4

    workdir = app4.config["PROCESSED_FOLDER"]
    workdir.mkdir(parents=True, exist_ok=True)
    phase4_token = uuid4().hex
    csv_path = workdir / f"phase4_ids_{phase4_token}.csv"
    csv_path.write_bytes(csv_bytes)

    df_ids = importer4.load_dataset(csv_path, delimiter=",")
    check(
        "leading-zero identifier preserved as text",
        df_ids["ID"].iloc[0] == "00123",
        f"first cell {df_ids['ID'].iloc[0]!r}, dtype {df_ids['ID'].dtype}",
    )
    id_token_path = workdir / f"phase4_id_token_{phase4_token}.csv"
    id_token_path.write_text("ID\n00123\nN/A\n00456\n", encoding="utf-8")
    df_id_token = importer4.load_dataset(id_token_path, delimiter=",")
    check(
        "leading-zero identifier survives a missing token in the same column",
        df_id_token["ID"].tolist() == ["00123", pd.NA, "00456"],
        f"got {df_id_token['ID'].tolist()}, dtype {df_id_token['ID'].dtype}",
    )
    check(
        "oversized integer stays exact (no rounding)",
        df_ids["Amount"].tolist() == [9007199254740993, 42, pd.NA],
        f"got {df_ids['Amount'].tolist()}, dtype {df_ids['Amount'].dtype}",
    )

    semi_path = workdir / f"phase4_semi_{phase4_token}.csv"
    semi_path.write_bytes(semi_bytes)
    df_latin = importer4.load_dataset(semi_path, delimiter=";", encoding="latin1")
    check(
        "latin1 encoding and semicolon delimiter honored",
        df_latin is not None
        and list(df_latin.columns) == ["Vornamen", "Sta_dtte"]
        and df_latin["Sta_dtte"].iloc[0] == "K\u00f6ln",
        f"cols={list(df_latin.columns) if df_latin is not None else None}",
    )
    df_latin_auto = importer4.load_dataset(semi_path, delimiter=";")
    check(
        "automatic latin1 fallback is used for both inferred and raw reads",
        df_latin_auto is not None and df_latin_auto["Sta_dtte"].iloc[0] == "K\u00f6ln",
    )

    # Lossless pickle round trip of supported types
    typed_df = pd.DataFrame({
        "ref": pd.array(["00123", "00124"], dtype="string"),
        "count": pd.array([5, None], dtype="Int64"),
        "flag": pd.array([True, False], dtype="boolean"),
        "day": pd.to_datetime(["2024-01-01", None]),
    })
    pkl_path = exporter4.export_dataset(typed_df, "phase4_roundtrip", file_format="pickle")
    back = importer4.load_dataset(pkl_path)
    check(
        "pickle intermediate preserves dtypes",
        back is not None
        and str(back["count"].dtype) == "Int64"
        and str(back["flag"].dtype) == "boolean"
        and str(back["day"].dtype).startswith("datetime64")
        and back["ref"].iloc[0] == "00123",
        f"dtypes={[str(back[c].dtype) for c in back.columns] if back is not None else None}",
    )
    pkl_path.unlink()

    # CSV/Excel exports generated directly from the current dataset
    csv_out = exporter4.export_dataset(typed_df, "phase4_export", file_format="csv", delimiter=";")
    csv_text = csv_out.read_bytes().decode("utf-8")
    check(
        "csv export keeps leading zeros and exact big values",
        "00123" in csv_text.splitlines()[1],
        csv_text[:120],
    )
    csv_out.unlink()

    xlsx_out = exporter4.export_dataset(typed_df, "phase4_export", file_format="excel")
    df_back = importer4.load_dataset(xlsx_out, sheet="Data")
    check(
        "excel export round trips values and accepts sheet selection",
        df_back is not None and df_back["count"].iloc[0] == 5,
    )
    check(
        "missing excel sheet surfaces workbook names (error path)",
        importer4.load_dataset(xlsx_out, sheet="Nope") is None,
    )
    xlsx_out.unlink()

    # Spreadsheet downloads neutralize formulas without changing workspace data.
    formula_values = [
        "=2+2",
        "+cmd",
        "-cmd",
        "@SUM(A1:A2)",
        "  =1+1",
        "\t=1+1",
        "\r=1+1",
        "\n=1+1",
        "'=already safe",
        "ordinary text",
    ]
    formula_df = pd.DataFrame({"=Formula heading": formula_values, "number": [-42] * 10})
    original_formula_df = formula_df.copy(deep=True)
    expected_safe_values = [
        "'=2+2",
        "'+cmd",
        "'-cmd",
        "'@SUM(A1:A2)",
        "'  =1+1",
        "'\t=1+1",
        "'\r=1+1",
        "'\n=1+1",
        "'=already safe",
        "ordinary text",
    ]

    safe_csv = exporter4.export_dataset(formula_df, "formula_safe", file_format="csv")
    csv_back = pd.read_csv(safe_csv, dtype=str, keep_default_na=False)
    check(
        "CSV export neutralizes formula cells and headings",
        list(csv_back.columns) == ["'=Formula heading", "number"]
        and csv_back.iloc[:, 0].tolist() == expected_safe_values,
        f"columns={list(csv_back.columns)}, values={csv_back.iloc[:, 0].tolist()}",
    )
    check("CSV export keeps negative numbers unprefixed", csv_back["number"].iloc[0] == "-42")
    safe_csv.unlink()

    safe_xlsx = exporter4.export_dataset(formula_df, "formula_safe", file_format="excel")
    workbook = load_workbook(safe_xlsx, data_only=False)
    worksheet = workbook["Data"]
    excel_values = [worksheet.cell(row=row, column=1).value for row in range(2, 12)]
    check(
        "Excel export stores formula-like values as text",
        worksheet.cell(row=1, column=1).value == "'=Formula heading"
        and all(value.startswith("'") for value in excel_values[:-1])
        and excel_values[-1] == "ordinary text"
        and all(worksheet.cell(row=row, column=1).data_type != "f" for row in range(1, 12)),
        f"heading={worksheet.cell(row=1, column=1).value!r}, values={excel_values}",
    )
    check(
        "Excel export keeps negative numbers numeric",
        worksheet.cell(row=2, column=2).value == -42
        and worksheet.cell(row=2, column=2).data_type == "n",
    )
    workbook.close()
    safe_xlsx.unlink()

    formula_pickle = exporter4.export_dataset(formula_df, "formula_lossless", file_format="pickle")
    check(
        "pickle snapshots preserve formula-like source text",
        pd.read_pickle(formula_pickle).equals(original_formula_df),
    )
    formula_pickle.unlink()
    check(
        "spreadsheet exports do not mutate the working dataset",
        formula_df.equals(original_formula_df),
    )

    csv_path.unlink()
    id_token_path.unlink()
    semi_path.unlink()

# --------------------------------------------------------------------------
# 6e. Phase 5 â€” predictable dates and numbers.
# --------------------------------------------------------------------------
df_dates = pd.DataFrame({"when": ["01/02/2024", "not-a-date"]})
df_month, _ = cleaning_service.convert_column_type(
    df_dates, "when", "date", "coerce_to_missing"
)
check(
    "date parse defaults to month-first (01/02 -> Jan 2)",
    pd.Timestamp("2024-01-02") == df_month["when"].iloc[0],
    f"got {df_month['when'].iloc[0]!r}",
)
df_day, rep_day = cleaning_service.convert_column_type(
    df_dates, "when", "date", "coerce_to_missing", dayfirst=True
)
check(
    "day-first setting is honored (01/02 -> Feb 1)",
    pd.Timestamp("2024-02-01") == df_day["when"].iloc[0],
    f"got {df_day['when'].iloc[0]!r}",
)
df_explicit, _ = cleaning_service.convert_column_type(
    pd.DataFrame({"when": ["31/01/2024"]}),
    "when",
    "date",
    date_format="%d/%m/%Y",
)
check(
    "explicit date format is honored",
    df_explicit["when"].iloc[0] == pd.Timestamp("2024-01-31"),
)
ambiguous_issues = analysis_service.detect_basic_issues(
    pd.DataFrame({"exam_date": ["01/02/2024", "03/04/2024"]})
)
check(
    "ambiguous dates are reported even when every value parses",
    any(issue["issue"] == "ambiguous_date" for issue in ambiguous_issues),
    str(ambiguous_issues),
)
bad_grouping, bad_grouping_report = cleaning_service.convert_column_type(
    pd.DataFrame({"amount": ["12,34"]}),
    "amount",
    "float",
    number_format="us",
)
check(
    "malformed numeric grouping is rejected",
    not bad_grouping_report.applied and bad_grouping["amount"].iloc[0] == "12,34",
)
check(
    "ambiguous date counted for analysis",
    analysis_service.count_ambiguous_dates(df_dates["when"][df_dates["when"] != "not-a-date"]) == 1,
)

df_eu = pd.DataFrame({"v": ["1.234,56", "12,5"]})
df_eu_conv, _ = cleaning_service.convert_column_type(
    df_eu, "v", "float", "coerce_to_missing", number_format="eu"
)
check(
    "european number format parsed (1.234,56 -> 1234.56)",
    df_eu_conv["v"].tolist() == [1234.56, 12.5],
    f"got {df_eu_conv['v'].tolist()}",
)
df_us = pd.DataFrame({"n": ["1,234"]})
df_us_conv, _ = cleaning_service.convert_column_type(
    df_us, "n", "integer", "coerce_to_missing", number_format="us"
)
check(
    "us thousands separator parsed (1,234 -> 1234)",
    df_us_conv["n"].tolist() == [1234],
    f"got {df_us_conv['n'].tolist()}",
)

df_tz = pd.DataFrame({"t": ["2024-01-01T10:00:00+05:00", "2024-01-02T10:00:00+05:00"]})
df_tz_conv, rep_tz = cleaning_service.convert_column_type(df_tz, "t", "date", "coerce_to_missing")
check(
    "timezone offsets preserved on date conversion",
    "+05:00" in str(df_tz_conv["t"].dtype),
    f"dtype {df_tz_conv['t'].dtype}",
)
check("uniform timezone not normalized", rep_tz.tz_normalized is False)

df_tz_mixed = pd.DataFrame({"t": ["2024-01-01T10:00:00+05:00", "2024-01-01T10:00:00+00:00"]})
df_tz_mixed_conv, rep_mixed = cleaning_service.convert_column_type(
    df_tz_mixed, "t", "date", "coerce_to_missing"
)
check(
    "mixed timezones are preserved unless UTC normalization is requested",
    rep_mixed.tz_normalized is False
    and str(df_tz_mixed_conv["t"].iloc[0].tzinfo) != str(df_tz_mixed_conv["t"].iloc[1].tzinfo),
    f"values={df_tz_mixed_conv['t'].tolist()}, report={rep_mixed.describe('t', 'date')}",
)
df_tz_utc, rep_utc = cleaning_service.convert_column_type(
    df_tz_mixed,
    "t",
    "date",
    "coerce_to_missing",
    timezone_policy="utc",
)
check(
    "mixed timezones normalize to UTC only when requested",
    rep_utc.tz_normalized is True
    and "UTC" in rep_utc.describe("t", "date")
    and str(df_tz_utc["t"].dtype).endswith("UTC]"),
    rep_utc.describe("t", "date"),
)

# --------------------------------------------------------------------------
# 6f. Phase 6 â€” validated cleaning plans.
# --------------------------------------------------------------------------
df_plan = pd.DataFrame({"a": [1, 2], "b": ["x ", "y"]})

errors = cleaning_service.validate_plan(
    df_plan, [{"type": "missing", "column": "zzz", "method": "remove"}]
)
check("plan rejects unknown column", errors and "does not exist" in errors[0], str(errors))

errors = cleaning_service.validate_plan(
    df_plan,
    [
        {"type": "drop_columns", "columns": ["a"]},
        {"type": "convert", "column": "a", "target_type": "float"},
    ],
)
check("plan rejects incompatible drop+use combination", errors, str(errors))

errors = cleaning_service.validate_plan(
    df_plan, [{"type": "case", "column": "b", "case_type": "weird"}]
)
check("plan rejects unknown parameter", errors and "case type" in errors[0], str(errors))

errors = cleaning_service.validate_plan(
    df_plan, [{"type": "convert", "column": "b", "target_type": "float", "invalid_policy": "nuked"}]
)
check("plan rejects unknown conversion policy", errors, str(errors))

errors = cleaning_service.validate_plan(df_plan, [{"type": "make_money"}])
check("plan rejects unknown operation type", errors, str(errors))

check(
    "valid plan passes with no errors",
    cleaning_service.validate_plan(
        df_plan, [{"type": "trim_whitespace"}, {"type": "duplicates", "mode": "last"}]
    )
    == [],
)

df_order = pd.DataFrame({"v": ["x ", " x ", "y"]})
_, order_hist = cleaning_service.apply_cleaning_operations(
    df_order, [{"type": "trim_whitespace"}, {"type": "duplicates", "mode": "first"}]
)
check(
    "trim runs before duplicate removal (near-duplicates caught)",
    any("Trimmed" in line for line in order_hist[:1])
    and any("1 duplicate" in line for line in order_hist),
    str(order_hist),
)

# Route-level: empty and duplicate rename targets are rejected.
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        ren_csrf = sess.get("csrf_token", "test-token")
    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "rename_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )
    columns = ["StudentID", "Name", "Age", "Gender", "Email", "Department", "CGPA", "EnrollmentDate"]
    res_dup = client.post(
        "/dataset/select-columns",
        data={
            "csrf_token": current_csrf(client),
            "selected_columns": columns,
            "rename_Name": "Zed",
            "rename_Department": "Zed",
        },
        follow_redirects=True,
    )
    check("duplicate rename target rejected", b"used more than once" in res_dup.data, "no flash seen")
    res_empty = client.post(
        "/dataset/select-columns",
        data={"csrf_token": current_csrf(client), "selected_columns": columns, "rename_Name": "   "},
        follow_redirects=True,
    )
    check("empty rename target rejected", b"cannot be empty" in res_empty.data, "no flash seen")

# --------------------------------------------------------------------------
# 6g. Phase 7 â€” preview before applying.
# --------------------------------------------------------------------------
df_prev = pd.DataFrame({"v": ["a", "b", None, "a"]})
preview_ops = [
    {"type": "trim_whitespace"},
    {"type": "missing", "column": "v", "method": "mode"},
    {"type": "duplicates", "mode": "first"},
]
preview_result = cleaning_service.preview_plan(df_prev, preview_ops)
check(
    # canonical order: trim (2) -> dedup (3) -> missing fill (6), so the
    # duplicate 'a' row is removed before the mode fill happens.
    "preview reports row counts",
    preview_result["rows_before"] == 4 and preview_result["rows_after"] == 3,
    str(preview_result),
)
check(
    "preview reports before/after examples",
    len(preview_result["examples"]) > 0
    and preview_result["examples"][0]["before"] != preview_result["examples"][0]["after"],
    str(preview_result["examples"]),
)

applied_df, _ = cleaning_service.apply_cleaning_operations(df_prev, preview_ops)
check(
    "applying an unchanged preview matches the previewed result",
    len(applied_df) == preview_result["rows_after"]
    and len(applied_df) != len(df_prev)
)

df_source = pd.DataFrame({"v": ["a", "b", None, "a"]})
cleaning_service.preview_plan(df_source, preview_ops)
check(
    "preview never mutates the working dataframe",
    # pandas may store None as NaN; compare missing semantics, not identity
    df_source["v"].isin(["a", "b"]).tolist() == [True, True, False, True],
    f"got {df_source['v'].tolist()}",
)

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        prev_csrf = sess.get("csrf_token", "test-token")
    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "preview_test.csv"), "delimiter": ",", "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )
    with client.session_transaction() as sess:
        original_path = Path(sess["upload_path"])
        original_bytes = original_path.read_bytes()
    prev_response = client.post(
        "/dataset/preview",
        data={"op_duplicates": "1", "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check("preview route renders", prev_response.status_code == 200 and b"Preview (nothing applied yet)" in prev_response.data)
    with client.session_transaction() as sess:
        stored_preview = sess.get("preview", {})
        token = stored_preview.get("token", "")
        rows_after = (stored_preview.get("result") or {}).get("rows_after")
    check("preview bound to parameters", token != "")
    # op_duplicates alone removes 1 exact duplicate; the second needs trimming
    check("preview computed removed rows", rows_after == 19, f"rows_after={rows_after}")
    check(
        "preview leaves the working file untouched",
        original_path.read_bytes() == original_bytes,
    )

    applied_response = client.post(
        "/clean",
        data={"op_duplicates": "1", "preview_token": token, "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check("fresh preview token applied", applied_response.status_code == 200)
    with client.session_transaction() as sess:
        preview_consumed = "preview" not in sess
    check("applied preview is consumed", preview_consumed)

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        stale_csrf = sess.get("csrf_token", "test-token")
    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "stale_preview_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )
    client.post(
        "/dataset/preview",
        data={"op_duplicates": "1", "csrf_token": current_csrf(client)},
    )
    with client.session_transaction() as sess:
        sess["preview"]["dataset"] = {"path": "stale"}  # simulate a dataset change
        stale_token = sess["preview"]["token"]
    stale_response = client.post(
        "/clean",
        data={"op_duplicates": "1", "preview_token": stale_token, "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check(
        "stale preview application rejected",
        b"out of date" in stale_response.data,
        "no stale rejection flash",
    )

# --------------------------------------------------------------------------
# 12. Phase 8 â€” undo & reproducible history.
# --------------------------------------------------------------------------
df_rep = pd.DataFrame({"a": [" x ", "y ", "  x  "], "c": [1, None, 3]})
recipe = [
    {"type": "trim_whitespace"},
    {"type": "missing", "column": "c", "method": "median"},
    {"type": "convert", "column": "c", "target_type": "integer"},
]
out_first, hist_first = cleaning_service.apply_cleaning_operations(
    df_rep, [dict(op) for op in recipe]
)
out_replay, hist_replay = cleaning_service.apply_cleaning_operations(
    df_rep, [dict(op) for op in recipe]
)
check(
    "replaying a recipe on the same input reproduces the result",
    list(out_first.columns) == list(out_replay.columns)
    and len(out_first) == len(out_replay)
    and [
        cleaning_service._values_match(a, b)
        for ra, rb in zip(out_first.to_numpy(dtype=object), out_replay.to_numpy(dtype=object))
        for a, b in zip(ra, rb)
    ].count(True)
    == len(out_first) * 2
    and hist_first == hist_replay,
)
check(
    "recipe produces the typed result",
    str(out_first["c"].dtype) == "Int64" and out_first["c"].tolist() == [1, 2, 3],
    f"dtype={out_first['c'].dtype}, values={out_first['c'].tolist()}",
)

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        undo_csrf = sess.get("csrf_token", "test-token")
    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "undo_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )

    clean_res = client.post(
        "/clean",
        data={"op_duplicates": "1", "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check("cleaned run recorded", clean_res.status_code == 200)

    audit = client.get("/download/audit")
    check(
        "audit download works with the recipe",
        audit.status_code == 200
        and audit.mimetype == "application/json"
        and b"duplicates" in audit.data,
    )
    with client.session_transaction() as sess:
        stored_ops = sess["audit_log"][-1]["operations"]
    check("recipe stores ordered operations with parameters", stored_ops == [{"type": "duplicates", "mode": "first"}], str(stored_ops))

    undo_res = client.post("/dataset/undo", data={"csrf_token": current_csrf(client)}, follow_redirects=True)
    check("undo flashes success", undo_res.status_code == 200)
    overview = client.get("/dataset")
    check("undo restores original 20 rows", b"20" in overview.data)

    undo_again = client.post("/dataset/undo", data={"csrf_token": current_csrf(client)}, follow_redirects=True)
    check("undo with empty stack reports nothing to undo", b"No cleaning runs" in undo_again.data)

# Undo restores a narrowed (column-selection) state, not just the original.
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        narrow_csrf = sess.get("csrf_token", "test-token")
    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "narrow_undo_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )
    columns = ["StudentID", "Name", "Age", "Gender", "Department", "CGPA", "EnrollmentDate"]
    client.post(
        "/dataset/select-columns",
        data={"csrf_token": current_csrf(client), "selected_columns": columns},
        follow_redirects=True,
    )
    client.post(
        "/clean",
        data={"op_trim": "1", "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    undo_res = client.post("/dataset/undo", data={"csrf_token": current_csrf(client)}, follow_redirects=True)
    overview = client.get("/dataset")
    check(
        "undo returns to the narrowed version (Email still dropped, 7 columns)",
        b"Email" not in overview.data and b"CGPA" in overview.data,
        "Email unexpectedly present after undo",
    )

# Session end removes version files from the undo stack.
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        end8_csrf = sess.get("csrf_token", "test-token")
    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "end8_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )
    client.post("/clean", data={"op_duplicates": "1", "csrf_token": current_csrf(client)}, follow_redirects=True)
    with client.session_transaction() as sess:
        stack_files = [Path(e["file"]) for e in sess.get("history_stack", [])]
    check("version files exist before end", stack_files and all(p.exists() for p in stack_files))
    client.post("/session/end", data={"csrf_token": current_csrf(client)})
    check("session end removes version files", all(not p.exists() for p in stack_files))


# --------------------------------------------------------------------------
# 7. Integration: Multi-step "Clean More" Workflow via Flask Test Client
# --------------------------------------------------------------------------

# CSRF token must come from the same client (server-side sessions), so seed it in-flow
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        csrf = sess.get("csrf_token", "test-token")

    with open(messy_csv, "rb") as fh:
        client.post("/upload", data={"file": (fh, "sample_messy_data.csv"), "csrf_token": current_csrf(client)})
    with client.session_transaction() as sess:
        reset_baseline = Path(sess["baseline_path"])

    # Pass 1: Trim whitespace & deduplicate
    res1 = client.post("/clean", data={"op_duplicates": "1", "op_trim": "1", "csrf_token": current_csrf(client)}, follow_redirects=True)
    check("pass 1 clean succeeds", res1.status_code == 200 and b"Removed 2 duplicate row(s)" in res1.data)

    # Inspect overview/analysis for "Clean More"
    res_analysis = client.get("/dataset/analysis")
    check("clean more loads cleaned state", res_analysis.status_code == 200)

    # Pass 2: Fill CGPA with median on top of pass 1
    res2 = client.post(
        "/clean",
        data={
            "misscol_1": "CGPA",
            "missing_1": "median",
            "csrf_token": current_csrf(client),
        },
        follow_redirects=True,
    )
    check("pass 2 clean succeeds", res2.status_code == 200)
    # Check that history accumulated both pass 1 and pass 2!
    check(
        "history accumulated pass 1 and pass 2",
        b"Removed 2 duplicate row(s)" in res2.data and b"median" in res2.data,
    )

    # Reset to original upload
    res_reset = client.post(
        "/dataset/reset",
        data={"csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check(
        "reset redirects to dataset overview",
        res_reset.status_code == 200 and b"Column Information" in res_reset.data,
    )
    # Overview after reset should have the original 20 rows
    check("reset restored original 20 rows", b"20" in res_reset.data)
    with client.session_transaction() as sess:
        check(
            "reset restores the lossless baseline snapshot",
            sess.get("current_path") == str(reset_baseline) and reset_baseline.exists(),
        )

# --------------------------------------------------------------------------
# 8. Explorer datatype selection persists into Cleaning
# --------------------------------------------------------------------------
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        type_csrf = sess.get("csrf_token", "test-token")

    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "type_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )

    columns = [
        "StudentID", "Name", "Age", "Gender", "Email", "Department",
        "CGPA", "EnrollmentDate",
    ]
    type_response = client.post(
        "/dataset/select-columns",
        data={
            "csrf_token": current_csrf(client),
            "selected_columns": columns,
            "type_StudentID": "Text",
        },
    )
    check("datatype selection redirects to Cleaning", type_response.status_code == 302)

    with client.session_transaction() as sess:
        saved_types = sess.get("column_type_overrides", {})
    check(
        "datatype selection is stored",
        saved_types.get("StudentID") == "string",
        f"stored overrides: {saved_types}",
    )

    type_page = client.get("/dataset/analysis").get_data(as_text=True)
    student_type_area = type_page[type_page.find("StudentID"):type_page.find("StudentID") + 500]
    check("selected datatype appears in Cleaning", "text" in student_type_area.lower())

    client.post("/dataset/reset", data={"csrf_token": current_csrf(client)})
    with client.session_transaction() as sess:
        check("reset clears datatype selections", "column_type_overrides" not in sess)

# --------------------------------------------------------------------------
# 9. Ending a session deletes its files and state
# --------------------------------------------------------------------------
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        end_csrf = sess.get("csrf_token", "test-token")

    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "end_session_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )

    client.post(
        "/clean",
        data={"op_trim": "1", "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    client.get("/download?format=excel")

    workspace_html = client.get("/dataset").data
    check(
        "active workspace offers end session",
        b'id="end-session-btn"' in workspace_html
        and b'id="end-session-modal"' in workspace_html,
    )

    with client.session_transaction() as sess:
        session_paths = {
            Path(path)
            for key in (
                "upload_path", "baseline_path", "current_path", "processed_path",
                "excel_path",
            )
            if (path := sess.get(key))
        }
    check("session files exist before ending", session_paths and all(path.exists() for path in session_paths))

    end_response = client.post(
        "/session/end",
        data={"csrf_token": current_csrf(client)},
        follow_redirects=False,
    )
    check(
        "end session redirects to fresh upload",
        end_response.status_code == 302 and end_response.headers["Location"].endswith("/upload"),
    )
    check("end session deletes all files", all(not path.exists() for path in session_paths))
    with client.session_transaction() as sess:
        check("end session clears progress", "original_filename" not in sess)

# --------------------------------------------------------------------------
# 10. Test Failed Upload File Cleanup
# --------------------------------------------------------------------------
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        failed_upload_csrf = sess.get("csrf_token", "test-token")
    # Upload corrupted file
    client.post(
        "/upload",
        data={"file": (io.BytesIO(b"\x00\x01broken"), "unparseable_test.csv"), "csrf_token": current_csrf(client)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    # Verify no unparseable_test.csv lingers in uploads
    broken_exists = (app.config["UPLOAD_FOLDER"] / "unparseable_test.csv").exists()
    check("failed upload unlinks file from disk", not broken_exists)

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        snapshot_csrf = sess.get("csrf_token", "test-token")
    client.post(
        "/upload",
        data={
            "file": (io.BytesIO(b"id,value\n1,kept\n"), "snapshot_kept.csv"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
    )
    with client.session_transaction() as sess:
        workspace_before_failure = sess.get("current_path")
        owned_before_failure = {
            Path(path) for path in sess.get("owned_paths", [])
        }

    failed_snapshot_name = f"snapshot_failure_{uuid4().hex}.csv"
    real_export = dataset_exporter.export_dataset

    def fail_snapshot(*args, **kwargs):
        raise OSError("simulated snapshot failure")

    dataset_exporter.export_dataset = fail_snapshot
    try:
        failed_snapshot_response = client.post(
            "/upload",
            data={
                "file": (io.BytesIO(b"id,value\n2,rejected\n"), failed_snapshot_name),
                "csrf_token": current_csrf(client),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
    finally:
        dataset_exporter.export_dataset = real_export

    with client.session_transaction() as sess:
        check(
            "snapshot failure keeps the active workspace",
            b"prepare this dataset" in failed_snapshot_response.data
            and sess.get("current_path") == workspace_before_failure
            and all(path.exists() for path in owned_before_failure),
        )
    check(
        "snapshot failure removes the rejected source",
        not (app.config["UPLOAD_FOLDER"] / failed_snapshot_name).exists(),
    )
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

# --------------------------------------------------------------------------
# 11. Route-level: dedup_mode and dedup_columns form fields are honored.
# --------------------------------------------------------------------------
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        dedup_csrf = sess.get("csrf_token", "test-token")
    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "dedup_mode_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )
    res_last = client.post(
        "/clean",
        data={"op_duplicates": "1", "dedup_mode": "last", "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check(
        "route honors keep-last mode with retained count",
        res_last.status_code == 200
        and b"kept the last occurrence" in res_last.data
        and b"row(s) retained" in res_last.data,
    )

    # A second pass exercising key-column selection end to end.
    columns = ["StudentID", "Name", "Age", "Gender", "Email", "Department", "CGPA", "EnrollmentDate"]
    res_subset = client.post(
        "/clean",
        data={"op_duplicates": "1", "dedup_mode": "first", "dedup_columns": ["Name", "Department"], "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check(
        "route honors selected key columns",
        b"matching keys (Name, Department)" in res_subset.data,
    )

# --------------------------------------------------------------------------
# 13. Repair coverage: atomic plans, dataset lifecycle and complete recipes.
# --------------------------------------------------------------------------
token_empty = pd.DataFrame({"empty": ["N/A", "--", " "], "keep": [1, 2, 3]})
token_empty_cleaned, token_empty_count = cleaning_service.remove_empty_columns(
    token_empty, ("n/a", "--")
)
check(
    "configured missing tokens make a column fully empty",
    token_empty_count == 1 and list(token_empty_cleaned.columns) == ["keep"],
)

invalid_custom_plan = [
    {"type": "trim_whitespace"},
    {"type": "missing", "column": "age", "method": "custom", "value": "abc"},
]
invalid_custom_errors = cleaning_service.validate_plan(
    pd.DataFrame({"age": [1.0, None], "name": [" a ", "b"]}),
    invalid_custom_plan,
)
check(
    "invalid custom fill rejects the whole plan before execution",
    bool(invalid_custom_errors) and "not a number" in invalid_custom_errors[0],
    str(invalid_custom_errors),
)

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        lifecycle_csrf = sess.get("csrf_token", "test-token")
    client.post(
        "/upload",
        data={
            "file": (io.BytesIO(b"age,name\n1, a \n, b \n"), "atomic_first.csv"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
    )
    with client.session_transaction() as sess:
        first_upload = Path(sess["upload_path"])
        first_baseline = Path(sess["baseline_path"])
        check(
            "accepted upload starts from a lossless baseline snapshot",
            first_upload.suffix == ".csv"
            and first_baseline.suffix == ".pkl"
            and first_upload.exists()
            and first_baseline.exists()
            and sess.get("current_path") == str(first_baseline)
            and "processed_path" not in sess,
        )
    rejected = client.post(
        "/clean",
        data={
            "op_trim": "1",
            "misscol_1": "age",
            "missing_1": "custom",
            "missingval_1": "abc",
            "csrf_token": current_csrf(client),
        },
        follow_redirects=True,
    )
    with client.session_transaction() as sess:
        check(
            "route leaves dataset untouched when plan validation fails",
            b"dataset was not changed" in rejected.data
            and sess.get("current_path") == str(first_baseline)
            and not sess.get("history_stack"),
        )

    client.post(
        "/clean",
        data={"op_trim": "1", "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    with client.session_transaction() as sess:
        first_owned = {Path(path) for path in sess.get("owned_paths", [])}
        check("first dataset has committed files", len(first_owned) >= 2)

    client.post(
        "/upload",
        data={
            "file": (io.BytesIO(b"new_id,value\n1,new\n"), "atomic_second.csv"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
    )
    with client.session_transaction() as sess:
        second_upload = Path(sess["upload_path"])
        second_baseline = Path(sess["baseline_path"])
        check(
            "new upload starts a clean workflow state",
            second_upload.exists()
            and second_baseline.exists()
            and second_baseline.suffix == ".pkl"
            and sess.get("current_path") == str(second_baseline)
            and not sess.get("history_stack")
            and not sess.get("audit_log")
            and "before_stats" not in sess
            and "processed_path" not in sess,
        )
    check(
        "new upload deletes every file owned by the previous dataset",
        all(not path.exists() for path in first_owned),
        str(first_owned),
    )
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        preview_dirty_csrf = sess.get("csrf_token", "test-token")
    client.post(
        "/upload",
        data={
            "file": (io.BytesIO(b"v\na\na\n"), "dirty_preview.csv"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
    )
    client.post(
        "/dataset/preview",
        data={"op_duplicates": "1", "csrf_token": current_csrf(client)},
    )
    with client.session_transaction() as sess:
        dirty_token = sess["preview"]["token"]
        dirty_original = sess["current_path"]
    dirty_response = client.post(
        "/clean",
        data={
            "op_trim": "1",
            "preview_token": dirty_token,
            "preview_dirty": "1",
            "csrf_token": current_csrf(client),
        },
        follow_redirects=True,
    )
    with client.session_transaction() as sess:
        check(
            "changed operations cannot be applied under an earlier preview",
            b"out of date" in dirty_response.data
            and sess.get("current_path") == dirty_original
            and not sess.get("history_stack"),
        )
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        version_csrf = sess.get("csrf_token", "test-token")
    client.post(
        "/upload",
        data={
            "file": (io.BytesIO(b"v\n a \n a \n"), "version_cleanup.csv"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
    )
    client.post("/clean", data={"op_trim": "1", "csrf_token": current_csrf(client)})
    client.post("/clean", data={"op_duplicates": "1", "csrf_token": current_csrf(client)})
    with client.session_transaction() as sess:
        version_paths = [Path(entry["file"]) for entry in sess["history_stack"]]
        source_for_undo = Path(sess["upload_path"])
        baseline_for_undo = Path(sess["baseline_path"])
    client.post("/dataset/undo", data={"csrf_token": current_csrf(client)})
    check(
        "undo immediately deletes the discarded version file",
        not version_paths[1].exists() and version_paths[0].exists(),
    )
    client.post("/dataset/undo", data={"csrf_token": current_csrf(client)})
    check(
        "undoing the first run restores the baseline snapshot",
        not version_paths[0].exists()
        and source_for_undo.exists()
        and baseline_for_undo.exists(),
    )
    with client.session_transaction() as sess:
        check(
            "first undo points at the baseline snapshot",
            sess.get("current_path") == str(baseline_for_undo),
        )
    check("reset endpoint rejects destructive GET", client.get("/dataset/reset").status_code == 405)
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        recipe_csrf = sess.get("csrf_token", "test-token")
    source_bytes = b"code,value,unused\n001,1,x\n002,2,y\n"
    client.post(
        "/upload",
        data={"file": (io.BytesIO(source_bytes), "selection_recipe.csv"), "csrf_token": current_csrf(client)},
        content_type="multipart/form-data",
    )
    with client.session_transaction() as sess:
        original_recipe_df = dataset_importer.load_dataset(Path(sess["upload_path"]), delimiter="auto")
    client.post(
        "/dataset/select-columns",
        data={
            "csrf_token": current_csrf(client),
            "selected_columns": ["code", "value"],
            "rename_value": "score",
            "type_value": "Text",
        },
    )
    with client.session_transaction() as sess:
        recipe_payload = [
            {"event": entry["event"], "operations": entry["operations"]}
            for entry in sess["history_stack"]
        ]
        selected_path = Path(sess["current_path"])
    replayed_selection = cleaning_service.replay_recipe(original_recipe_df, recipe_payload)
    selected_df = dataset_importer.load_dataset(selected_path)
    check(
        "column selection recipe reproduces selection, rename and datatype",
        list(replayed_selection.columns) == ["code", "score"]
        and replayed_selection.equals(selected_df)
        and str(replayed_selection["score"].dtype) == "string",
        f"recipe={recipe_payload}",
    )
    audit_payload = client.get("/download/audit").get_json()
    check(
        "downloaded audit exposes the active replayable selection recipe",
        audit_payload["recipe"] == recipe_payload,
        str(audit_payload["recipe"]),
    )

    client.post("/clean", data={"op_trim": "1", "csrf_token": current_csrf(client)})
    csv_download = client.get("/download?format=csv")
    with client.session_transaction() as sess:
        csv_download_path = Path(sess["csv_path"])
        csv_is_owned = str(csv_download_path) in sess.get("owned_paths", [])
    check(
        "clean CSV download is exported uniquely and tracked for cleanup",
        csv_download.status_code == 200 and csv_download_path.exists() and csv_is_owned,
    )
    csv_download.close()
    client.post("/session/end", data={"csrf_token": current_csrf(client)})
    check("ending the session deletes generated CSV downloads", not csv_download_path.exists())

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        sheet_csrf = sess.get("csrf_token", "test-token")
    workbook = io.BytesIO()
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        pd.DataFrame({"first_col": [1]}).to_excel(writer, sheet_name="First", index=False)
        pd.DataFrame({"second_col": [2]}).to_excel(writer, sheet_name="Second", index=False)
    workbook.seek(0)
    client.post(
        "/upload",
        data={
            "file": (workbook, "selected_sheet.xlsx"),
            "sheet": "Second",
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
    )
    sheet_response = client.post(
        "/dataset/select-columns",
        data={"selected_columns": ["second_col"], "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check(
        "selected Excel worksheet is retained through column selection",
        sheet_response.status_code == 200
        and b"second_col" in sheet_response.data
        and b"first_col" not in sheet_response.data,
    )
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

# --------------------------------------------------------------------------
# 14. Case-only cleaning keeps unresolved category mapping visible.
# --------------------------------------------------------------------------
with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        category_csrf = sess.get("csrf_token", "test-token")
    category_csv = (
        "Gender,Row\nF,1\nFEM,2\nFem,3\nM,4\nMALE,5\nMale,6\n"
        "Extremely Long Category Name,7\nextremely long category name,8\n"
    ).encode("utf-8")
    client.post(
        "/upload",
        data={
            "file": (io.BytesIO(category_csv), "category_review.csv"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
    )
    initial_cleaning_page = client.get("/dataset/analysis").get_data(as_text=True)
    check(
        "category source values render fully inside the bulk mapper",
        'title="Extremely Long Category Name">Extremely Long Category Name</span>'
        in initial_cleaning_page
        and "category-map-source-check" in initial_cleaning_page
        and "category-map-target" in initial_cleaning_page
        and "add-category-map" in initial_cleaning_page
        and "Standard name..." not in initial_cleaning_page,
    )
    client.post(
        "/clean",
        data={
            "textcol_1": "Gender",
            "case_1": "lower",
            "csrf_token": current_csrf(client),
        },
        follow_redirects=True,
    )
    clean_more_page = client.get("/dataset/analysis").get_data(as_text=True)
    with client.session_transaction() as sess:
        pending_categories = sess.get("pending_category_columns", [])
    check(
        "case-only cleaning keeps category mapping review visible",
        pending_categories == ["Gender"]
        and "Case was standardized, but these values still need mapping review." in clean_more_page
        and "Map values in groups" in clean_more_page,
        f"pending={pending_categories}",
    )
    check(
        "type conversion editor starts with one compact editor",
        clean_more_page.count('name="convert_column"') == 0
        and 'id="conversion-column"' in clean_more_page
        and 'novalidate' in clean_more_page
        and 'id="conversion-list"' in clean_more_page
        and "Add conversion" in clean_more_page
        and "Add Row" not in clean_more_page,
    )
    client.post(
        "/clean",
        data={
            "textcol_1": "Gender",
            "mapvar_1_1_1": "f",
            "map_1_1_1": "female",
            "mapvar_1_1_2": "fem",
            "map_1_1_2": "female",
            "mapvar_1_2_1": "m",
            "map_1_2_1": "male",
            "mapvar_1_2_2": "male",
            "map_1_2_2": "male",
            "csrf_token": current_csrf(client),
        },
    )
    with client.session_transaction() as sess:
        mapped_category_df = dataset_importer.load_dataset(Path(sess["current_path"]))
        check(
            "an explicit category mapping clears the pending review marker",
            "pending_category_columns" not in sess,
        )
    check(
        "multiple source labels map to grouped standard labels in one pass",
        mapped_category_df["Gender"].tolist()[:6]
        == ["female", "female", "female", "male", "male", "male"],
        str(mapped_category_df["Gender"].tolist()),
    )
    client.post("/dataset/undo", data={"csrf_token": current_csrf(client)})
    with client.session_transaction() as sess:
        check(
            "undo restores the pending category review state",
            sess.get("pending_category_columns") == ["Gender"],
        )
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

# --------------------------------------------------------------------------
# 15. Phase 9 â€” broader analysis findings (counts + examples, read-only).
# --------------------------------------------------------------------------
df_mixed = pd.DataFrame({"score": ["10", "12", "not-a-number", "15"]})
mixed = analysis_service.detect_mixed_types(df_mixed)
check(
    "mixed types detected with count and examples",
    mixed.get("score", {}).get("count") == 1
    and mixed["score"]["examples"] == ["not-a-number"]
    and mixed["score"]["dominant"] == "number",
    str(mixed),
)

df_keys = pd.DataFrame({"customer_id": [1, 2, 2, 3], "name": ["a", "b", "c", "d"]})
key_findings = analysis_service.detect_duplicate_keys(df_keys)
check(
    "duplicate keys detected on a key-looking column",
    key_findings["customer_id"]["count"] == 1 and key_findings["customer_id"]["examples"] == ["2"],
    str(key_findings),
)
check(
    "non-key columns are not reported as duplicate keys",
    "name" not in key_findings,
    str(key_findings),
)

df_cats = pd.DataFrame({"department": ["Sales", "Sales", "HR", "HR", "HR", "Salse"]})
category_findings = analysis_service.detect_invalid_categories(df_cats)
check(
    "one-off invalid category detected",
    category_findings["department"]["count"] == 1
    and category_findings["department"]["examples"] == ["Salse"],
    str(category_findings),
)

df_out = pd.DataFrame({"amount": [10, 11, 12, 13, 100]})
outliers = analysis_service.detect_outliers(df_out)
check(
    "numeric outliers flagged by the 1.5xIQR rule with bounds and examples",
    outliers["amount"]["count"] == 1
    and outliers["amount"]["multiplier"] == 1.5
    and outliers["amount"]["examples"] == ["100"]
    and outliers["amount"]["lower"] < outliers["amount"]["upper"],
    str(outliers),
)
check(
    "outlier detection never changes the data",
    df_out["amount"].tolist() == [10, 11, 12, 13, 100],
    str(df_out["amount"].tolist()),
)

analysis_full = analysis_service.analyze_dataset(df_cats)
check(
    "full analysis exposes anomalies in its suggestions",
    any(s["id"] == "anomalies" for s in analysis_full["suggestions"])
    and analysis_full["invalid_categories"]["department"]["count"] == 1,
    str([s["id"] for s in analysis_full["suggestions"]]),
)

# User-defined range violations are reported with counts and examples.
range_report = cleaning_service.evaluate_validation_rule(
    pd.DataFrame({"age": [5, 20, 99]}),
    {"type": "validate", "column": "age", "rule": "range", "minimum": 10, "maximum": 50},
)
check(
    "user-defined range violations counted with examples",
    range_report["count"] == 2 and range_report["examples"] == ["5", "99"],
    str(range_report),
)

# --------------------------------------------------------------------------
# 16. Phase 9 â€” new transformations (validation, preview and undo are generic).
# --------------------------------------------------------------------------
df_norm = pd.DataFrame({"note": ["a\u00a0\u00a0b", "  c  ", "d"]})
norm_df, norm_count = cleaning_service.normalize_text(df_norm)
check(
    "unicode and whitespace normalized",
    norm_df["note"].tolist() == ["a b", "c", "d"] and norm_count == 2,
    f"{norm_df['note'].tolist()} / {norm_count}",
)

regex_df, regex_count = cleaning_service.regex_replace(
    pd.DataFrame({"code": ["AB-12", "CD-34"]}), "code", r"^[A-Z]{2}-", ""
)
check(
    "regex replacement applies and counts cells",
    regex_df["code"].tolist() == ["12", "34"] and regex_count == 2,
    str(regex_df["code"].tolist()),
)

split_df, split_rows = cleaning_service.split_column(
    pd.DataFrame({"path": ["a/b", "c/d", None]}), "path", "/", ["left", "right"]
)
check(
    "column split into named columns and source dropped",
    split_rows == 2
    and "path" not in split_df.columns
    and split_df["left"].tolist()[:2] == ["a", "c"],
    f"{list(split_df.columns)} {split_df['left'].tolist()}",
)

merged_df, merged_rows = cleaning_service.merge_columns(
    pd.DataFrame({"first": ["Ann", "Bob"], "last": ["Lee", None]}),
    ["first", "last"],
    " ",
    "full",
    True,
)
check(
    "columns merged, blanks skipped",
    merged_rows == 2
    and merged_df["full"].tolist() == ["Ann Lee", "Bob"]
    and "first" not in merged_df.columns,
    str(merged_df["full"].tolist()),
)

df_group = pd.DataFrame({
    "region": ["x", "x", "y", "y"],
    "amount": [10.0, None, None, None],
})
group_df, group_report = cleaning_service.fill_by_group(df_group, ["region"], "amount", "mean")
check(
    "group-based filling uses only each group's values",
    group_df["amount"].tolist()[0:2] == [10.0, 10.0]
    and group_report["filled"] == 1
    and pd.isna(group_df["amount"].iloc[2]),
    f"{group_df['amount'].tolist()} {group_report}",
)
check(
    "unfillable groups are reported, not guessed",
    group_report["unfilled"] == 2 and group_report["unfilled_groups"] == ["y"],
    str(group_report),
)

# Validation, preview and undo for the new operations through the service.
new_ops = [
    {"type": "normalize_text", "columns": []},
    {"type": "regex_replace", "column": "name", "pattern": "^J", "replacement": "Y"},
    {"type": "split_column", "column": "code", "separator": "-", "new_columns": ["code_left", "code_right"]},
    {"type": "merge_columns", "columns": ["name", "region"], "separator": "-", "new_column": "who", "drop_source": False},
    {"type": "group_fill", "column": "amount", "group_by": ["region"], "method": "mean"},
    {"type": "validate", "column": "amount", "rule": "range", "minimum": 0, "maximum": 100, "action": "report"},
]
df_new = pd.DataFrame({
    "name": [" Jane ", "John", "Ana"],
    "code": ["a-1", "b-2", "c-3"],
    "region": ["x", "x", "y"],
    "amount": [10.0, None, 300.0],
})
check(
    "new operation plan validates cleanly",
    cleaning_service.validate_plan(df_new, [dict(op) for op in new_ops]) == [],
    str(cleaning_service.validate_plan(df_new, [dict(op) for op in new_ops])),
)
preview_new = cleaning_service.preview_plan(df_new, [dict(op) for op in new_ops])
check(
    "preview covers the new operations",
    len(preview_new["steps"]) == len(new_ops)
    and any(step["label"].startswith("split_column") for step in preview_new["steps"])
    and len(preview_new["examples"]) > 0,
    str(preview_new["steps"]),
)
new_result, new_history = cleaning_service.apply_cleaning_operations(
    df_new, [dict(op) for op in new_ops]
)
check(
    "new operations run and describe themselves",
    all(
        any(keyword in line for line in new_history)
        for keyword in ("Normalized Unicode", "Replaced text matching", "Split", "Merged", "Filled 1 missing", "Validation 'range'")
    ),
    str(new_history),
)

# Invalid new-operation parameters are rejected before anything runs.
check(
    "plan rejects an uncompileable regex",
    cleaning_service.validate_plan(
        df_new, [{"type": "regex_replace", "column": "name", "pattern": "(", "replacement": ""}]
    ),
)
check(
    "plan rejects split into an existing column name",
    cleaning_service.validate_plan(
        df_new,
        [{"type": "split_column", "column": "code", "separator": "-", "new_columns": ["name"]}],
    ),
)
check(
    "plan rejects a single-column merge",
    cleaning_service.validate_plan(
        df_new, [{"type": "merge_columns", "columns": ["name"], "separator": "-", "new_column": "m"}]
    ),
)
check(
    "plan rejects an unknown group-fill method",
    cleaning_service.validate_plan(
        df_new, [{"type": "group_fill", "column": "amount", "group_by": ["region"], "method": "average"}]
    ),
)
check(
    "plan rejects a non-numeric column for a mean group fill",
    cleaning_service.validate_plan(
        df_new, [{"type": "group_fill", "column": "name", "group_by": ["region"], "method": "mean"}]
    ),
)

# Category mappings report the values they left untouched.
sample_df, sample_history = cleaning_service.apply_cleaning_operations(
    pd.DataFrame({"gender": ["m", "male", "f", "unknown"]}),
    [{"type": "categories", "column": "gender", "mapping": {"m": "male"}}],
)
check(
    "unmapped categories are reported after a mapping",
    sample_df["gender"].tolist()[:2] == ["male", "male"]
    and any("had no mapping" in line for line in sample_history),
    str(sample_history),
)

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        tools_csrf = sess.get("csrf_token", "test-token")
    tools_csv = "name,code,region,amount\n Jane ,a-1,x,10\nJohn,b-2,x,\nAna,c-3,y,300\n".encode("utf-8")
    client.post(
        "/upload",
        data={"file": (io.BytesIO(tools_csv), "tools_test.csv"), "csrf_token": current_csrf(client)},
        content_type="multipart/form-data",
    )
    tools_form = {
        "csrf_token": current_csrf(client),
        "op_normalize": "1",
        "regex_column": "name",
        "regex_pattern": "^J",
        "regex_replacement": "Y",
        "split_column": "code",
        "split_separator": "-",
        "split_new_columns": "code_left, code_right",
        "merge_columns": ["name", "region"],
        "merge_new_column": "who",
        "merge_separator": "-",
        "merge_keep": "1",
        "groupfill_column": "amount",
        "groupfill_by": ["region"],
        "groupfill_method": "mean",
        "validate_column": "amount",
        "validate_rule": "range",
        "validate_min": "0",
        "validate_max": "100",
        "validate_action": "report",
    }
    preview_page = client.post("/dataset/preview", data=tools_form, follow_redirects=True)
    check(
        "cleaning page renders the new findings and tools",
        b"Validation &amp; Anomalies" in preview_page.data
        and b"Text &amp; Column Tools" in preview_page.data
        and b"Group-Based Filling" in preview_page.data
        and b"Validation Rules" in preview_page.data,
        "new panels missing from the cleaning page",
    )
    check(
        "preview works for the new operations",
        b"Preview (nothing applied yet)" in preview_page.data,
    )
    applied = client.post("/clean", data=tools_form, follow_redirects=True)
    check(
        "new operations apply through the form",
        applied.status_code == 200
        and b"Normalized Unicode" in applied.data
        and b"Filled 1 missing" in applied.data
        and b"Validation &#39;range&#39;" in applied.data,
        "expected history lines missing",
    )
    with client.session_transaction() as sess:
        applied_df = dataset_importer.load_dataset(Path(sess["current_path"]))
    check(
        "applied result has the split and merged columns",
        list(applied_df.columns) == ["name", "region", "amount", "code_left", "code_right", "who"]
        and applied_df["code_left"].tolist() == ["a", "b", "c"],
        str(list(applied_df.columns)),
    )
    client.post("/dataset/undo", data={"csrf_token": current_csrf(client)}, follow_redirects=True)
    with client.session_transaction() as sess:
        undone_df = dataset_importer.load_dataset(Path(sess["current_path"]))
    check(
        "undo reverses the new operations",
        "code" in undone_df.columns and "who" not in undone_df.columns,
        str(list(undone_df.columns)),
    )
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

# --------------------------------------------------------------------------
# 16b. Phase 9 â€” user-defined validation rules and their actions.
# --------------------------------------------------------------------------
df_rules = pd.DataFrame({
    "status": ["ok", "bad", "ok", "ok", "ok"],
    "amount": [10, 11, 12, 13, 500],
})
allowed_report = cleaning_service.evaluate_validation_rule(
    df_rules,
    {"type": "validate", "column": "status", "rule": "allowed", "values": ["ok"]},
)
check(
    "user-defined allowed categories flag invalid ones with examples",
    allowed_report["count"] == 1 and allowed_report["examples"] == ["bad"],
    str(allowed_report),
)
outlier_report = cleaning_service.evaluate_validation_rule(
    df_rules, {"type": "validate", "column": "amount", "rule": "outlier"}
)
check(
    "outlier rule flags via the 1.5xIQR bounds with examples",
    outlier_report["count"] == 1
    and outlier_report["examples"] == ["500"]
    and "1.5" in outlier_report["details"],
    str(outlier_report),
)
reported_df, _ = cleaning_service.apply_validation_rule(
    df_rules, {"type": "validate", "column": "amount", "rule": "outlier", "action": "report"}
)
check(
    "flagging outliers never changes the values",
    reported_df["amount"].tolist() == [10, 11, 12, 13, 500],
    str(reported_df["amount"].tolist()),
)
blanked_df, _ = cleaning_service.apply_validation_rule(
    df_rules,
    {
        "type": "validate",
        "column": "status",
        "rule": "allowed",
        "values": ["ok"],
        "action": "set_missing",
    },
)
check(
    "set-missing action blanks exactly the violating values",
    pd.isna(blanked_df["status"].iloc[1]) and blanked_df["status"].iloc[0] == "ok",
    str(blanked_df["status"].tolist()),
)
removed_df, _ = cleaning_service.apply_validation_rule(
    df_rules, {"type": "validate", "column": "amount", "rule": "outlier", "action": "remove"}
)
check(
    "remove action drops only the violating rows",
    removed_df["amount"].tolist() == [10, 11, 12, 13],
    str(removed_df["amount"].tolist()),
)

check(
    "plan rejects an unknown validation rule",
    cleaning_service.validate_plan(
        df_rules, [{"type": "validate", "column": "status", "rule": "sortof"}]
    ),
)
check(
    "plan rejects a range with no limits and an inverted range",
    cleaning_service.validate_plan(
        df_rules, [{"type": "validate", "column": "amount", "rule": "range"}]
    )
    and cleaning_service.validate_plan(
        df_rules,
        [{"type": "validate", "column": "amount", "rule": "range", "minimum": 10, "maximum": 1}],
    ),
)
check(
    "plan rejects an uncompileable rule pattern and an empty allowed list",
    cleaning_service.validate_plan(
        df_rules, [{"type": "validate", "column": "status", "rule": "pattern", "pattern": "("}]
    )
    and cleaning_service.validate_plan(
        df_rules, [{"type": "validate", "column": "status", "rule": "allowed", "values": []}]
    ),
)
check(
    "plan rejects normalizing a column that does not exist",
    cleaning_service.validate_plan(df_rules, [{"type": "normalize_text", "columns": ["nope"]}]),
)

# --------------------------------------------------------------------------
# 17. Phase 10 â€” configurable limits and actionable parsing errors.
# --------------------------------------------------------------------------
check(
    "row limit produces an actionable message",
    "over the limit of"
    in dataset_limit_error(pd.DataFrame({"a": range(Config.MAX_ROWS + 1)})),
)

original_max_columns = Config.MAX_COLUMNS
original_max_cells = Config.MAX_CELLS
Config.MAX_COLUMNS = 2
Config.MAX_CELLS = 3
try:
    check(
        "column limit is enforced",
        "3 columns, over the limit of 2"
        in dataset_limit_error(pd.DataFrame({"a": [1], "b": [2], "c": [3]})),
    )
    check(
        "cell limit is enforced",
        "6 cells, over the limit of 3"
        in dataset_limit_error(pd.DataFrame({"a": [1, 2, 3], "b": [4, 5, 6]})),
    )
finally:
    Config.MAX_COLUMNS = original_max_columns
    Config.MAX_CELLS = original_max_cells

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        limit_csrf = sess.get("csrf_token", "test-token")
    original_max_rows = Config.MAX_ROWS
    Config.MAX_ROWS = 5
    try:
        with open(messy_csv, "rb") as fh:
            rejected = client.post(
                "/upload",
                data={"file": (fh, "too_big_rows.csv"), "csrf_token": current_csrf(client)},
                content_type="multipart/form-data",
                follow_redirects=True,
            )
        check("oversized upload rejected with the limit numbers", b"over the limit of" in rejected.data)
        check(
            "rejected oversized upload is removed from disk",
            not (app.config["UPLOAD_FOLDER"] / "too_big_rows.csv").exists(),
        )
    finally:
        Config.MAX_ROWS = original_max_rows

    broken_csv = b'a,b\n"unclosed,1\n'
    ragged_upload = client.post(
        "/upload",
        data={
            "file": (io.BytesIO(broken_csv), "unreadable_upload.csv"),
            "delimiter": ",",
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    check(
        "unreadable upload explains the failure",
        b"do not line up as CSV" in ragged_upload.data,
        "no parse reason flashed",
    )
    check(
        "unreadable upload is removed from disk",
        not (app.config["UPLOAD_FOLDER"] / "unreadable_upload.csv").exists(),
    )

ragged_path = Path(app.config["PROCESSED_FOLDER"]) / "phase10_ragged.csv"
ragged_path.write_bytes(broken_csv)
parsed, reason = dataset_importer.load_dataset_with_error(ragged_path, delimiter=",")
check(
    "parser errors become an actionable sentence",
    parsed is None and "line up as CSV" in reason,
    reason,
)
check(
    "load_dataset still returns None quietly for bad files",
    dataset_importer.load_dataset(ragged_path, delimiter=",") is None,
)
ragged_path.unlink()

# --------------------------------------------------------------------------
# 18. Phase 10 ? safe, report-first recipe review and application.
# --------------------------------------------------------------------------
recipe_payload = {
    "format": cleaning_service.RECIPE_FORMAT,
    "version": cleaning_service.RECIPE_VERSION,
    "recipe": [
        {
            "event": "select_columns",
            "operations": [{
                "type": "select_columns",
                "columns": ["name", "score"],
                "renames": {"name": "student"},
                "conversions": {"score": "float"},
            }],
        },
        {
            "event": "clean",
            "operations": [
                {"type": "trim_whitespace"},
                {
                    "type": "validate",
                    "column": "score",
                    "rule": "range",
                    "minimum": 0,
                    "maximum": 100,
                    "action": "report",
                    "tokens": [],
                },
            ],
        },
    ],
}
normalized_recipe = cleaning_service.normalize_recipe_payload(recipe_payload)
recipe_source = pd.DataFrame({
    "name": [" Alice ", "Bob"],
    "score": [10, 120],
    "extra": ["x", "y"],
})
recipe_result, recipe_report, recipe_types, recipe_history = cleaning_service.inspect_recipe(
    recipe_source, normalized_recipe
)
check(
    "recipe dry-run applies selection, rename and conversion on a copy",
    list(recipe_result.columns) == ["student", "score"]
    and recipe_result["student"].tolist() == ["Alice", "Bob"]
    and str(recipe_result["score"].dtype) == "Float64"
    and recipe_source["name"].iloc[0] == " Alice ",
    f"{recipe_result.dtypes} / {recipe_result.to_dict('list')}",
)
check(
    "recipe review reports extra columns and report-only violations",
    recipe_report["runs"] == 2
    and any("extra dataset column" in warning for warning in recipe_report["warnings"])
    and any("values will stay unchanged" in warning for warning in recipe_report["warnings"])
    and recipe_types == {"score": "float"},
    str(recipe_report),
)
check(
    "recipe replay produces the reviewed result",
    cleaning_service.replay_recipe(recipe_source, normalized_recipe).equals(recipe_result),
)

for label, bad_payload in (
    ("wrong recipe version", {**recipe_payload, "version": 2}),
    ("unknown recipe event", {**recipe_payload, "recipe": [{"event": "mystery", "operations": [{"type": "trim_whitespace"}]}]}),
    ("malformed recipe operations", {**recipe_payload, "recipe": [{"event": "clean", "operations": ["trim"]}]}),
):
    try:
        cleaning_service.normalize_recipe_payload(bad_payload)
        rejected = False
    except ValueError:
        rejected = True
    check(f"{label} is rejected", rejected)

try:
    cleaning_service.inspect_recipe(
        recipe_source.drop(columns=["score"]), normalized_recipe
    )
    missing_rejected = False
except ValueError as exc:
    missing_rejected = "score" in str(exc)
check("recipe requires exact column names", missing_rejected)

try:
    cleaning_service.inspect_recipe(
        recipe_source,
        [{"event": "clean", "operations": [{"type": "mystery"}]}],
    )
    unknown_operation_rejected = False
except ValueError:
    unknown_operation_rejected = True
check("recipe rejects unknown cleaning operations", unknown_operation_rejected)

collision_recipe = [{
    "event": "select_columns",
    "operations": [{
        "type": "select_columns",
        "columns": ["name", "score"],
        "renames": {"name": "value", "score": "value"},
        "conversions": {},
    }],
}]
try:
    cleaning_service.inspect_recipe(recipe_source, collision_recipe)
    collision_rejected = False
except ValueError:
    collision_rejected = True
check("recipe rejects column-name collisions", collision_rejected)

multi_rules = [
    {"type": "validate", "column": "score", "rule": "range", "minimum": 0, "maximum": 100, "action": "report", "tokens": []},
    {"type": "validate", "column": "name", "rule": "required", "action": "report", "tokens": []},
]
reported, _ = cleaning_service.apply_cleaning_operations(recipe_source, multi_rules)
check("multiple report-first validation rules do not mutate data", reported.equals(recipe_source))
check(
    "non-finite ranges, risky patterns, one-output splits and repeated merge sources are rejected",
    bool(cleaning_service.validate_plan(recipe_source, [{"type": "validate", "column": "score", "rule": "range", "minimum": "nan"}]))
    and bool(cleaning_service.validate_plan(recipe_source, [{"type": "regex_replace", "column": "name", "pattern": "(a+)+$", "replacement": ""}]))
    and bool(cleaning_service.validate_plan(recipe_source, [{"type": "split_column", "column": "name", "separator": " ", "new_columns": ["only"]}]))
    and bool(cleaning_service.validate_plan(recipe_source, [{"type": "split_column", "column": "name", "separator": " ", "new_columns": ["name", "last"]}]))
    and bool(cleaning_service.validate_plan(recipe_source, [{"type": "merge_columns", "columns": ["name", "name"], "separator": " ", "new_column": "merged"}])),
)
zero_iqr = cleaning_service.evaluate_validation_rule(
    pd.DataFrame({"x": [5, 5, 5, 5, 5]}),
    {"type": "validate", "column": "x", "rule": "outlier"},
)
check("zero-IQR validation matches anomaly detection", zero_iqr["count"] == 0)

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        recipe_csrf = sess.get("csrf_token", "test-token")
    recipe_csv = b"name,score,extra\n Alice ,10,x\nBob,120,y\n"
    client.post(
        "/upload",
        data={"file": (io.BytesIO(recipe_csv), "recipe_target.csv"), "csrf_token": current_csrf(client)},
        content_type="multipart/form-data",
    )
    with client.session_transaction() as sess:
        before_recipe_path = sess["current_path"]
    multi_preview = client.post(
        "/dataset/preview",
        data={
            "csrf_token": current_csrf(client),
            "validation_rule_json": [json.dumps(rule) for rule in multi_rules],
        },
        follow_redirects=True,
    )
    with client.session_transaction() as sess:
        parsed_rules = sess["preview"]["operations"]
    check(
        "multi-rule builder payload reaches one report-first preview",
        multi_preview.status_code == 200
        and len(parsed_rules) == 2
        and all(rule["action"] == "report" for rule in parsed_rules),
        str(parsed_rules),
    )
    oversized_recipe = client.post(
        "/recipe/inspect",
        data={
            "recipe_file": (io.BytesIO(b" " * (Config.MAX_RECIPE_BYTES + 1)), "large.json"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    check("oversized recipe file is rejected", b"1 MB or smaller" in oversized_recipe.data)
    with client.session_transaction() as sess:
        check("invalid recipe review leaves the dataset untouched", sess["current_path"] == before_recipe_path)
    reviewed = client.post(
        "/recipe/inspect",
        data={
            "recipe_file": (io.BytesIO(json.dumps(recipe_payload).encode()), "recipe.json"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    check(
        "recipe route shows a compatibility review before applying",
        b"Apply reviewed recipe" in reviewed.data and b"extra dataset column" in reviewed.data,
        reviewed.data[:300].decode(errors="ignore"),
    )
    with client.session_transaction() as sess:
        recipe_token = sess["recipe_preview"]["token"]
        check("recipe review does not change the working dataset", sess["current_path"] == before_recipe_path)
    applied = client.post(
        "/recipe/apply",
        data={"recipe_token": recipe_token, "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check("reviewed recipe applies as one result", b"Applied 2 reviewed recipe run(s)" in applied.data)
    with client.session_transaction() as sess:
        recipe_df = dataset_importer.load_dataset(Path(sess["current_path"]))
        recipe_entry = sess["history_stack"][-1]
    check(
        "recipe application stores one undoable imported-recipe event",
        list(recipe_df.columns) == ["student", "score"]
        and recipe_entry["event"] == "recipe_import"
        and recipe_entry["source_recipe"] == normalized_recipe,
        str(recipe_entry),
    )
    audit = client.get("/download/audit").get_json()
    check("imported recipe remains replayable in the audit export", audit["recipe"] == normalized_recipe)
    client.post("/dataset/undo", data={"csrf_token": current_csrf(client)})
    with client.session_transaction() as sess:
        undone_recipe_df = dataset_importer.load_dataset(Path(sess["current_path"]), delimiter="auto")
    check("one undo reverses the entire imported recipe", list(undone_recipe_df.columns) == ["name", "score", "extra"])

    client.post(
        "/recipe/inspect",
        data={
            "recipe_file": (io.BytesIO(json.dumps(recipe_payload).encode()), "recipe.json"),
            "csrf_token": current_csrf(client),
        },
        content_type="multipart/form-data",
    )
    with client.session_transaction() as sess:
        stale_token = sess["recipe_preview"]["token"]
        stale_path = Path(sess["current_path"])
    with stale_path.open("ab") as changed_file:
        changed_file.write(b"\n")
    stale = client.post(
        "/recipe/apply",
        data={"recipe_token": stale_token, "csrf_token": current_csrf(client)},
        follow_redirects=True,
    )
    check("stale recipe review is rejected", b"review is out of date" in stale.data)
    with client.session_transaction() as sess:
        check("stale recipe rejection is atomic", sess["current_path"] == str(stale_path))
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

# --------------------------------------------------------------------------
# 20. Phase 10 â€” session file expiry and isolated output paths.
# --------------------------------------------------------------------------
with app.app_context():
    processed_root = Path(app.config["PROCESSED_FOLDER"])
    stale_file = processed_root / "phase10_stale.csv"
    stale_file.write_text("x", encoding="utf-8")
    os.utime(stale_file, (time.time() - 600, time.time() - 600))
    fresh_file = processed_root / "phase10_fresh.csv"
    fresh_file.write_text("y", encoding="utf-8")
    removed = file_store.sweep_stale_files(60)
    check(
        "abandoned files are expired, fresh ones kept",
        removed >= 1 and not stale_file.exists() and fresh_file.exists(),
        f"removed={removed}",
    )
    check(
        "a long TTL keeps files",
        file_store.sweep_stale_files(10 ** 9) == 0 and fresh_file.exists(),
    )
    fresh_file.unlink()

with app.test_client() as client:
    client.get("/upload")
    with client.session_transaction() as sess:
        ns_csrf = sess.get("csrf_token", "test-token")
    with open(messy_csv, "rb") as fh:
        client.post(
            "/upload",
            data={"file": (fh, "namespace_test.csv"), "csrf_token": current_csrf(client)},
            content_type="multipart/form-data",
        )
    client.post("/clean", data={"op_trim": "1", "csrf_token": current_csrf(client)}, follow_redirects=True)
    with client.session_transaction() as sess:
        namespace = sess.get("output_namespace")
        output_path = Path(sess["processed_path"])
    check(
        "outputs live in a per-session folder",
        bool(namespace) and output_path.parent.name == namespace and output_path.exists(),
        f"namespace={namespace} path={output_path}",
    )
    client.post("/session/end", data={"csrf_token": current_csrf(client)})

print()
if failures:
    print(f"FAILED: {len(failures)} test(s) failed: {failures}")
    sys.exit(1)
print("All regression tests passed successfully!")
